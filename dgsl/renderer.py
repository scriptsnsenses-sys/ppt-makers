"""DGSL renderer — resolved IR -> native editable .pptx via python-pptx.

Coordinate system: DGSL positions and sizes are plain numbers on a
1280 x 720 canvas, mapped onto the 16:9 slide (13.33 x 7.5 in).
Everything produced is a genuine PowerPoint shape / text frame / table /
chart, so the file stays fully editable.
"""

import io
import logging

from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

try:
    import requests
except Exception:
    requests = None

from .errors import DGSLError

log = logging.getLogger("ppt-maker.dgsl")

DGSL_W, DGSL_H = 1280.0, 720.0
SLIDE_W_IN, SLIDE_H_IN = 13.33, 7.5

MAX_IMAGE_BYTES = 6 * 1024 * 1024
MAX_PPTX_BYTES = 25 * 1024 * 1024

SHAPE_MAP = {
    "rect": "RECTANGLE",
    "rounded_rect": "ROUNDED_RECTANGLE",
    "circle": "OVAL",
    "ellipse": "OVAL",
    "triangle": "ISOSCELES_TRIANGLE",
    "diamond": "DIAMOND",
    "star": "STAR_5_POINT",
    "hexagon": "HEXAGON",
    "pentagon": "PENTAGON",
    "chevron": "CHEVRON",
    "arrow": "RIGHT_ARROW",
    "callout": "RECTANGULAR_CALLOUT",
}

ALIGN_MAP = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
             "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
VALIGN_MAP = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
              "bottom": MSO_ANCHOR.BOTTOM}

CHART_MAP = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _xu(u):
    return Inches(SLIDE_W_IN * float(u) / DGSL_W)


def _yu(u):
    return Inches(SLIDE_H_IN * float(u) / DGSL_H)


def parse_color(value, line=None):
    if not isinstance(value, str):
        raise DGSLError("colors must be quoted hex strings like "
                        '"#7C3AED"', line=line)
    s = value.strip().lstrip("#")
    if len(s) == 3:
        s = "".join(c * 2 for c in s)
    if len(s) != 6 or any(c not in "0123456789abcdefABCDEF" for c in s):
        raise DGSLError(f"bad color {value!r} — use hex like "
                        '"#7C3AED"', line=line)
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def _num(value, what, line=None, allow_none=True):
    if value is None and allow_none:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise DGSLError(f"{what} must be a number", line=line)
    return float(value)


def _str(value, what, line=None):
    if not isinstance(value, str):
        raise DGSLError(f"{what} must be a quoted string", line=line)
    return value


def _bool(value, what, line=None):
    if not isinstance(value, bool):
        raise DGSLError(f"{what} must be true or false", line=line)
    return value


class Ctx:
    """Render context: collects deduped warnings."""

    def __init__(self, warnings):
        self.warnings = warnings
        self._seen = set(warnings)

    def warn(self, message, line=None):
        text = f"line {line}: {message}" if line is not None else message
        if text not in self._seen:
            self._seen.add(text)
            self.warnings.append(text)


def _download_image(url):
    if requests is None:
        return None
    if not isinstance(url, str):
        return None
    url = url.strip()
    if not (url.startswith("http://") or url.startswith("https://")):
        return None
    try:
        with requests.get(url, timeout=8, stream=True,
                           headers={"User-Agent": "ppt-maker-dgsl/1.0"}) as r:
            r.raise_for_status()
            ctype = (r.headers.get("Content-Type") or "").lower()
            if ctype and not ctype.startswith("image/") and \
                    "octet" not in ctype:
                return None
            buf = io.BytesIO()
            total = 0
            for chunk in r.iter_content(chunk_size=64 * 1024):
                if not chunk:
                    continue
                total += len(chunk)
                if total > MAX_IMAGE_BYTES:
                    return None
                buf.write(chunk)
            data = buf.getvalue()
            return data if len(data) >= 128 else None
    except Exception as exc:
        log.info("DGSL image download failed (%s): %s", url, exc)
        return None


def _set_alt_text(shape, alt):
    try:
        el = shape._element
        for attr in ("_nvSpPr", "_nvPicPr", "_nvCxnSpPr",
                     "nvGraphicFramePr", "_nvGrpSpPr"):
            holder = getattr(el, attr, None)
            if holder is not None and hasattr(holder, "cNvPr"):
                holder.cNvPr.set("descr", alt)
                return
    except Exception:
        pass


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def render_pptx(ir):
    """Render a compiled DGSL IR object and return the .pptx bytes."""
    from pptx import Presentation
    warnings = list(ir.warnings)
    ctx = Ctx(warnings)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    _apply_meta(prs, ir.meta)

    for slide in ir.slides:
        _render_slide(prs, slide, ctx)

    buf = io.BytesIO()
    prs.save(buf)
    data = buf.getvalue()
    if len(data) > MAX_PPTX_BYTES:
        raise DGSLError("generated presentation exceeds 25 MB — "
                        "use fewer or smaller images")
    ir.warnings = warnings
    return data


def _apply_meta(prs, meta):
    core = prs.core_properties
    try:
        if meta.get("title"):
            core.title = str(meta["title"])[:255]
        if meta.get("author"):
            core.author = str(meta["author"])[:255]
        if meta.get("subject"):
            core.subject = str(meta["subject"])[:255]
        if meta.get("keywords"):
            kw = meta["keywords"]
            core.keywords = ", ".join(str(k) for k in kw) \
                if isinstance(kw, list) else str(kw)[:255]
        if meta.get("language"):
            core.language = str(meta["language"])[:50]
        core.comments = "Generated with DGSL v1 (DG PPT Maker)"
    except Exception:
        pass


def _render_slide(prs, slide, ctx):
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    if slide.name:
        try:
            s.name = slide.name[:100]
        except Exception:
            pass
    _render_background(s, slide.background, ctx, slide.line)
    elements = sorted(slide.elements,
                      key=lambda e: _zorder(e))
    for el in elements:
        if not el["props"].get("visible", True):
            continue
        _render_element(s, el, ctx)
    if slide.notes:
        try:
            s.notes_slide.notes_text_frame.text = slide.notes
        except Exception:
            pass


def _zorder(el):
    z = el["props"].get("z", 0)
    return z if isinstance(z, (int, float)) and not isinstance(z, bool) \
        else 0


# ---------------------------------------------------------------------------
# Backgrounds
# ---------------------------------------------------------------------------

def _render_background(slide, bg, ctx, line):
    if bg is None:
        _solid_bg(slide, RGBColor(0xFF, 0xFF, 0xFF))
        return
    kind = bg.get("kind")
    if kind == "solid":
        _solid_bg(slide, parse_color(bg.get("color", "#FFFFFF"), line))
    elif kind == "gradient":
        fill = slide.background.fill
        try:
            fill.gradient()
            stops = fill.gradient_stops
            stops[0].color.rgb = parse_color(bg["from"], line)
            stops[0].position = 0.0
            stops[1].color.rgb = parse_color(bg["to"], line)
            stops[1].position = 1.0
            try:
                fill.gradient_angle = float(bg.get("angle", 90))
            except Exception:
                pass
        except Exception as exc:
            raise DGSLError(f"could not render gradient: {exc}", line=line)
    elif kind == "image":
        data = _download_image(bg.get("src", ""))
        if data:
            try:
                pic = slide.shapes.add_picture(
                    io.BytesIO(data), Inches(0), Inches(0),
                    width=Inches(SLIDE_W_IN), height=Inches(SLIDE_H_IN))
                _send_to_back(slide, pic)
            except Exception:
                ctx.warn("background image could not be embedded", line)
                _solid_bg(slide, RGBColor(0xFF, 0xFF, 0xFF))
        else:
            ctx.warn(f"background image {bg.get('src', '')!r} is not "
                     f"reachable — use an http(s) URL to embed it", line)
            _solid_bg(slide, RGBColor(0xFF, 0xFF, 0xFF))
    else:
        _solid_bg(slide, RGBColor(0xFF, 0xFF, 0xFF))


def _solid_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _send_to_back(slide, shape):
    try:
        spTree = slide.shapes._spTree
        spTree.remove(shape._element)
        spTree.insert(2, shape._element)
    except Exception:
        pass


# ---------------------------------------------------------------------------
# Elements
# ---------------------------------------------------------------------------

def _render_element(slide, el, ctx):
    kind = el["kind"]
    if kind == "text":
        _render_text(slide, el, ctx)
    elif kind in SHAPE_MAP:
        _render_shape(slide, el, ctx)
    elif kind == "line":
        _render_line(slide, el, ctx)
    elif kind == "image":
        _render_image(slide, el, ctx)
    elif kind == "table":
        _render_table(slide, el, ctx)
    elif kind == "chart":
        _render_chart(slide, el, ctx)
    elif kind == "group":
        _render_group(slide, el, ctx)
    elif kind == "button":
        _render_button(slide, el, ctx)
    else:  # video / audio / icon / connector placeholders
        _render_placeholder(slide, el, ctx)


def _geom(el, defaults, line):
    p = el["props"]
    x = _num(p.get("x", defaults[0]), "x", line, allow_none=False)
    y = _num(p.get("y", defaults[1]), "y", line, allow_none=False)
    w = _num(p.get("width", defaults[2]), "width", line, allow_none=False)
    h = _num(p.get("height", defaults[3]), "height", line, allow_none=False)
    if w <= 0 or h <= 0:
        raise DGSLError("width and height must be positive numbers",
                        line=line)
    return _xu(x), _yu(y), _xu(w), _yu(h)


def _apply_common(shape, el, ctx):
    p = el["props"]
    line = el.get("line")
    if el.get("id"):
        try:
            shape.name = str(el["id"])[:100]
        except Exception:
            pass
    if el.get("alt"):
        _set_alt_text(shape, str(el["alt"]))
    if "rotation" in p:
        try:
            shape.rotation = _num(p["rotation"], "rotation", line,
                                  allow_none=False)
        except DGSLError:
            raise
        except Exception:
            ctx.warn("rotation is not supported for this element", line)
    if "opacity" in p:
        ctx.warn("opacity is parsed but transparency is not rendered "
                 "in DGSL v1", line)


def _apply_textbox_extras(text_frame, shape, el, ctx):
    """Fill + border behind a text box."""
    p = el["props"]
    line = el.get("line")
    if p.get("background") is not None:
        try:
            fill = shape.fill
            fill.solid()
            fill.fore_color.rgb = parse_color(p["background"], line)
        except DGSLError:
            raise
        except Exception:
            ctx.warn("text background could not be applied", line)
    if el.get("border"):
        _apply_border(shape, el["border"], ctx, line, kind="text")
    elif "border" in p:
        ctx.warn("'border' on text needs a border { ... } block", line)


def _apply_border(shape, border, ctx, line, kind="shape"):
    color = border.get("color")
    width = border.get("width")
    if color is not None:
        try:
            shape.line.color.rgb = parse_color(color, line)
        except DGSLError:
            raise
        except Exception:
            ctx.warn("border color could not be applied", line)
    if width is not None:
        w = _num(width, "border width", line, allow_none=False)
        if w < 0:
            raise DGSLError("border width cannot be negative", line=line)
        try:
            shape.line.width = Pt(w)
        except Exception:
            ctx.warn("border width could not be applied", line)
    if border.get("radius") is not None and kind == "text":
        ctx.warn("rounded corners on text boxes are not rendered "
                 "in DGSL v1", line)
    for key in border:
        if key not in ("color", "width", "radius"):
            ctx.warn(f"unknown border property '{key}' ignored", line)


# -- text ------------------------------------------------------------------

def _render_text(slide, el, ctx, box=None):
    p = el["props"]
    line = el.get("line")
    _check_props(el, {"x", "y", "width", "height", "font", "size", "color",
                      "bold", "italic", "underline", "strikethrough",
                      "align", "valign", "line_spacing", "letter_spacing",
                      "paragraph_spacing", "background", "rotation",
                      "opacity", "z", "visible", "id", "link", "language"},
                 ctx)
    left, top, width, height = _geom(el, (80, 80, 960, 120), line)
    if box is None:
        box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True

    if "valign" in p:
        va = p["valign"]
        if va not in VALIGN_MAP:
            ctx.warn(f"valign must be top, middle, or bottom "
                     f"(found {va!r})", line)
        else:
            tf.vertical_anchor = VALIGN_MAP[va]

    runs = el["runs"] or [{"text": el.get("text") or ""}]
    parent_font = {"font": p.get("font"), "size": p.get("size"),
                   "color": p.get("color"), "bold": p.get("bold"),
                   "italic": p.get("italic"), "underline": p.get("underline")}
    if p.get("strikethrough"):
        ctx.warn("strikethrough is parsed but not rendered in DGSL v1",
                 line)
    first = True
    for run in runs:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _style_paragraph(para, p, ctx, line)
        prun = para.add_run()
        prun.text = str(run.get("text", ""))
        merged = dict(parent_font)
        merged.update({k: v for k, v in run.items()
                       if v is not None and k != "text"})
        _style_run(prun, merged, ctx, line)
        if el.get("link"):
            try:
                prun.hyperlink.address = str(el["link"])
            except Exception:
                ctx.warn("hyperlink could not be applied", line)

    _apply_textbox_extras(tf, box, el, ctx)
    _apply_common(box, el, ctx)


def _style_paragraph(para, p, ctx, line):
    if "align" in p:
        al = p["align"]
        if al not in ALIGN_MAP:
            ctx.warn(f"align must be left, center, right, or justify "
                     f"(found {al!r})", line)
        else:
            para.alignment = ALIGN_MAP[al]
    if "line_spacing" in p:
        para.line_spacing = _num(p["line_spacing"], "line_spacing", line,
                                 allow_none=False)
    if "paragraph_spacing" in p:
        para.space_after = Pt(_num(p["paragraph_spacing"],
                                   "paragraph_spacing", line,
                                   allow_none=False))
    if "letter_spacing" in p:
        ctx.warn("letter_spacing is parsed but not rendered in DGSL v1",
                 line)


def _style_run(run, fmt, ctx, line):
    font = run.font
    if fmt.get("font") is not None:
        font.name = str(fmt["font"])
    if fmt.get("size") is not None:
        font.size = Pt(_num(fmt["size"], "size", line, allow_none=False))
    if fmt.get("color") is not None:
        font.color.rgb = parse_color(fmt["color"], line)
    for key in ("bold", "italic", "underline"):
        if fmt.get(key) is not None:
            setattr(font, key, _bool(fmt[key], key, line))
    if fmt.get("strikethrough"):
        ctx.warn("strikethrough is parsed but not rendered in DGSL v1",
                 line)


# -- shapes -----------------------------------------------------------------

def _render_shape(slide, el, ctx):
    kind = el["kind"]
    line = el.get("line")
    _check_props(el, {"x", "y", "width", "height", "fill", "radius",
                      "rotation", "opacity", "z", "visible", "id", "alt",
                      "link", "language"}, ctx)
    left, top, width, height = _geom(el, (80, 80, 400, 240), line)
    shape_name = SHAPE_MAP[kind]
    try:
        mso = getattr(MSO_SHAPE, shape_name)
    except AttributeError:
        ctx.warn(f"shape '{kind}' is not available — "
                 f"rendered as a rectangle", line)
        mso = MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(mso, left, top, width, height)

    p = el["props"]
    if p.get("fill") is not None:
        try:
            shape.fill.solid()
            shape.fill.fore_color.rgb = parse_color(p["fill"], line)
        except DGSLError:
            raise
        except Exception:
            ctx.warn("shape fill could not be applied", line)
    else:
        try:
            shape.fill.background()
        except Exception:
            pass

    if el.get("border"):
        _apply_border(shape, el["border"], ctx, line)
    else:
        try:
            shape.line.fill.background()
        except Exception:
            pass

    if kind == "rounded_rect" or p.get("radius") is not None:
        radius = p.get("radius", 16)
        radius = _num(radius, "radius", line, allow_none=False)
        if radius < 0:
            raise DGSLError("radius cannot be negative", line=line)
        if kind != "rounded_rect":
            ctx.warn("'radius' only rounds rounded_rect shapes", line)
        else:
            try:
                span = min(float(width.emu), float(height.emu))
                shape.adjustments[0] = max(
                    0.0, min(0.5, float(radius) / max(span, 1e-6) * 360))
            except Exception:
                pass

    if el.get("shadow"):
        pass  # warned at compile time; shapes stay flat in v1.
    try:
        shape.shadow.inherit = False
    except Exception:
        pass

    if el.get("link"):
        try:
            shape.click_action.hyperlink.address = str(el["link"])
        except Exception:
            ctx.warn("hyperlink could not be applied", line)
    _apply_common(shape, el, ctx)


# -- lines ------------------------------------------------------------------

def _render_line(slide, el, ctx):
    p = el["props"]
    line_no = el.get("line")
    _check_props(el, {"x1", "y1", "x2", "y2", "color", "width", "start",
                      "end", "z", "visible", "id", "rotation"}, ctx)
    x1 = _num(p.get("x1", 80), "x1", line_no, allow_none=False)
    y1 = _num(p.get("y1", 360), "y1", line_no, allow_none=False)
    x2 = _num(p.get("x2", 1200), "x2", line_no, allow_none=False)
    y2 = _num(p.get("y2", 360), "y2", line_no, allow_none=False)
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      _xu(x1), _yu(y1), _xu(x2), _yu(y2))
    if p.get("color") is not None:
        try:
            conn.line.color.rgb = parse_color(p["color"], line_no)
        except DGSLError:
            raise
        except Exception:
            pass
    if p.get("width") is not None:
        conn.line.width = Pt(_num(p["width"], "width", line_no,
                                  allow_none=False))
    if p.get("start") not in (None, "none") or \
            p.get("end") not in (None, "none"):
        ctx.warn("line arrowheads are parsed but not rendered in DGSL "
                 "v1", line_no)
    _apply_common(conn, el, ctx)


# -- images -----------------------------------------------------------------

def _render_image(slide, el, ctx):
    p = el["props"]
    line_no = el.get("line")
    _check_props(el, {"x", "y", "width", "height", "fit", "rotation",
                      "opacity", "z", "visible", "id", "alt", "link",
                      "language"}, ctx)
    src = el.get("src", "")
    left, top, width, height = _geom(el, (80, 80, 640, 360), line_no)
    if p.get("fit") not in (None, "contain", "cover"):
        ctx.warn(f"fit must be contain or cover (found "
                 f"{p.get('fit')!r})", line_no)
    data = _download_image(src)
    if data:
        try:
            pic = slide.shapes.add_picture(io.BytesIO(data), left, top,
                                           width=width, height=height)
            if el.get("border"):
                _apply_border(pic, el["border"], ctx, line_no)
            if el.get("link"):
                try:
                    pic.click_action.hyperlink.address = str(el["link"])
                except Exception:
                    pass
            _apply_common(pic, el, ctx)
            return
        except Exception as exc:
            log.info("DGSL image embed failed: %s", exc)
    label = el.get("alt") or src or "image"
    ctx.warn(f"image {src!r} is not reachable — placeholder rendered "
             f"(use an http(s) URL to embed it)", line_no)
    _placeholder_box(slide, left, top, width, height,
                     "Image placeholder", str(label)[:80])


def _placeholder_box(slide, left, top, width, height, title, sub=""):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   left, top, width, height)
    try:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xF1, 0xF5, 0xF9)
        shape.line.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        shape.line.width = Pt(1.5)
    except Exception:
        pass
    tf = shape.text_frame
    tf.word_wrap = True
    try:
        from pptx.enum.text import MSO_ANCHOR as _A
        tf.vertical_anchor = _A.MIDDLE
    except Exception:
        pass
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = title
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    if sub:
        para2 = tf.add_paragraph()
        para2.alignment = PP_ALIGN.CENTER
        run2 = para2.add_run()
        run2.text = sub
        run2.font.size = Pt(11)
        run2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    return shape


def _render_placeholder(slide, el, ctx):
    p = el["props"]
    line_no = el.get("line")
    left, top, width, height = _geom(el, (80, 80, 640, 360), line_no)
    label = {"video": "Video", "audio": "Audio", "icon": "Icon",
             "connector": "Connector"}.get(el["kind"], el["kind"])
    sub = el.get("src") or ""
    shape = _placeholder_box(slide, left, top, width, height, label, sub)
    _apply_common(shape, el, ctx)


# -- tables -----------------------------------------------------------------

def _render_table(slide, el, ctx):
    line_no = el.get("line")
    _check_props(el, {"x", "y", "width", "height", "z", "visible", "id",
                      "alt", "font", "size"}, ctx)
    columns = el["columns"]
    rows = el["rows"]
    if not columns and not rows:
        raise DGSLError("table needs columns [...] and at least one "
                        "row [...]", line=line_no)
    body = [list(map(str, r)) for r in rows]
    header = list(map(str, columns)) if columns else None
    n_cols = max([len(header or [])] + [len(r) for r in body] + [1])
    n_rows = len(body) + (1 if header else 0)
    if n_rows == 0:
        raise DGSLError("table needs at least one row [...]", line=line_no)

    left, top, width, height = _geom(el, (80, 160, 1120, 400), line_no)
    graphic = slide.shapes.add_table(n_rows, n_cols, left, top, width,
                                     height)
    try:
        graphic.name = (el.get("id") or "table")[:100]
    except Exception:
        pass
    if el.get("alt"):
        _set_alt_text(graphic, str(el["alt"]))
    tbl = graphic.table
    try:
        tbl.first_row = bool(header)
        tbl.horz_banding = False
    except Exception:
        pass

    col_w = int(width.emu // max(n_cols, 1))
    for j in range(n_cols):
        try:
            tbl.columns[j].width = col_w
        except Exception:
            pass

    header_cfg = el.get("header") or {}
    base_font = el["props"].get("font")
    base_size = el["props"].get("size")
    borders_cfg = el.get("borders_cfg") or {}

    def fill_row(i, values, is_header):
        for j in range(n_cols):
            cell = tbl.cell(i, j)
            text = values[j] if j < len(values) else ""
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            run = para.add_run()
            run.text = text
            if is_header:
                if header_cfg.get("fill"):
                    _cell_fill(cell, header_cfg["fill"], line_no)
                run.font.bold = bool(header_cfg.get("bold", True))
                if header_cfg.get("color"):
                    run.font.color.rgb = parse_color(header_cfg["color"],
                                                     line_no)
                else:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    if not header_cfg.get("fill"):
                        _cell_fill(cell, "#312E81", line_no)
                if header_cfg.get("size") is not None:
                    run.font.size = Pt(_num(header_cfg["size"], "size",
                                            line_no, allow_none=False))
                elif base_size is not None:
                    run.font.size = Pt(_num(base_size, "size", line_no,
                                            allow_none=False))
                if header_cfg.get("font") or base_font:
                    run.font.name = str(header_cfg.get("font") or base_font)
            else:
                run.font.bold = False
                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                run.font.size = Pt(_num(base_size, "size", line_no,
                                        allow_none=False)
                                   if base_size is not None else 14)
                if base_font:
                    run.font.name = str(base_font)
            if borders_cfg:
                _cell_borders(cell, borders_cfg, line_no, ctx)

    if header:
        fill_row(0, header, True)
        for i, row in enumerate(body, start=1):
            fill_row(i, row, False)
    else:
        for i, row in enumerate(body):
            fill_row(i, row, False)


def _cell_fill(cell, color, line_no):
    try:
        cell.fill.solid()
        cell.fill.fore_color.rgb = parse_color(color, line_no)
    except DGSLError:
        raise
    except Exception:
        pass


def _cell_borders(cell, cfg, line_no, ctx):
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    color = cfg.get("color", "#CBD5E1")
    width = cfg.get("width", 1)
    try:
        hexcolor = "%02X%02X%02X" % parse_color(color, line_no)
    except DGSLError:
        raise
    w = _num(width, "border width", line_no, allow_none=False)
    try:
        tcPr = cell._tc.get_or_add_tcPr()
        for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            for old in tcPr.findall(qn(edge)):
                tcPr.remove(old)
            ln = OxmlElement(edge)
            ln.set("w", str(max(1, int(float(w) * 12700))))
            fill = OxmlElement("a:solidFill")
            srgb = OxmlElement("a:srgbClr")
            srgb.set("val", hexcolor)
            fill.append(srgb)
            ln.append(fill)
            tcPr.append(ln)
    except Exception:
        ctx.warn("table borders could not be applied", line_no)


# -- charts -----------------------------------------------------------------

def _render_chart(slide, el, ctx):
    from pptx.chart.data import ChartData
    line_no = el.get("line")
    _check_props(el, {"x", "y", "width", "height", "title", "color", "z",
                      "visible", "id", "alt"}, ctx)
    labels = [str(v) for v in el["labels"]]
    values = []
    for v in el["values"]:
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            raise DGSLError("chart values must be numbers", line=line_no)
        values.append(float(v))
    if not labels or not values:
        raise DGSLError("chart needs labels [...] and values [...]",
                        line=line_no)
    if len(labels) != len(values):
        raise DGSLError(f"chart needs the same number of labels "
                        f"({len(labels)}) and values ({len(values)})",
                        line=line_no)

    left, top, width, height = _geom(el, (80, 160, 1120, 480), line_no)
    data = ChartData()
    data.categories = labels
    data.add_series("", tuple(values))
    try:
        graphic = slide.shapes.add_chart(CHART_MAP[el["chart_type"]], left,
                                         top, width, height, data)
    except Exception as exc:
        raise DGSLError(f"could not build chart: {exc}", line=line_no)
    try:
        graphic.name = (el.get("id") or "chart")[:100]
    except Exception:
        pass
    if el.get("alt"):
        _set_alt_text(graphic, str(el["alt"]))
    chart = graphic.chart
    try:
        chart.style = 2
    except Exception:
        pass
    title = el["props"].get("title")
    if title is not None:
        try:
            chart.has_title = True
            chart.chart_title.text_frame.text = str(title)
        except Exception:
            pass
    try:
        chart.has_legend = False
    except Exception:
        pass
    color = el["props"].get("color", "#4F46E5")
    try:
        series = chart.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = parse_color(color, line_no)
    except DGSLError:
        raise
    except Exception:
        pass
    try:
        chart.value_axis.has_major_gridlines = False
    except Exception:
        pass


# -- groups & buttons --------------------------------------------------------

def _render_group(slide, el, ctx):
    p = el["props"]
    line_no = el.get("line")
    for key in ("x", "y", "width", "height", "rotation"):
        if key in p:
            ctx.warn(f"group '{key}' transforms are not applied in DGSL "
                     f"v1 — children keep their own geometry", line_no)
            break
    for child in sorted(el["children"], key=_zorder):
        if not child["props"].get("visible", True):
            continue
        _render_element(slide, child, ctx)


def _render_button(slide, el, ctx):
    line_no = el.get("line")
    left, top, width, height = _geom(el, (80, 80, 320, 96), line_no)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top,
                                   width, height)
    try:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        shape.line.fill.background()
    except Exception:
        pass
    tf = shape.text_frame
    tf.word_wrap = True
    try:
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    except Exception:
        pass
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = str(el.get("text") or "Button")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if el.get("link"):
        try:
            run.hyperlink.address = str(el["link"])
        except Exception:
            pass
    if el.get("target"):
        ctx.warn(f"button navigation to slide {el['target']!r} is not "
                 f"interactive in DGSL v1", line_no)
    _apply_common(shape, el, ctx)


# -- prop checking -----------------------------------------------------------

def _check_props(el, allowed, ctx):
    line = el.get("line")
    for key in el["props"]:
        if key not in allowed:
            ctx.warn(f"unknown '{el['kind']}' property '{key}' ignored",
                     line)
